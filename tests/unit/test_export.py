from datetime import date
from io import BytesIO
from types import SimpleNamespace
from openpyxl import load_workbook
from app.services.export_service import COLUMNS, build_sid_workbook

def test_exact_sid_columns_and_text_ids():
    m=SimpleNamespace(no_urut_kk=1,status_hubungan='Kepala Keluarga',nik='0012345678901234',nama_lengkap='DATA FIKTIF',jenis_kelamin='Laki-laki',tempat_lahir='TEST',tanggal_lahir=date(2000,1,2),agama='ISLAM',pendidikan='SMA',jenis_pekerjaan='PELAJAR',status_perkawinan='BELUM KAWIN',kewarganegaraan='WNI',no_paspor=None,no_kitas_kitap=None,nama_ayah='AYAH',nama_ibu='IBU',golongan_darah='O'); r=SimpleNamespace(no_kk='0012345678901234',nama_kepala_keluarga='DATA FIKTIF',alamat='ALAMAT UJI',rt='2',rw='1',kode_pos='51111',dusun='DUSUN',desa='DESA',kecamatan='KECAMATAN',kabupaten='KABUPATEN',provinsi='JAWA TENGAH',members=[m]); data,rows=build_sid_workbook([SimpleNamespace(kk_record=r)]); ws=load_workbook(BytesIO(data)).active; assert [c.value for c in ws[1]]==COLUMNS; assert ws['A2'].value=='0012345678901234'; assert ws['N2'].value=='0012345678901234'; assert ws['A2'].number_format=='@'; assert ws['N2'].number_format=='@'; assert ws['D2'].value=='002'; assert ws['E2'].value=='001'; assert rows==1
