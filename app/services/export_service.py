from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
COLUMNS=['no_kk','nama_kepala_keluarga','alamat','rt','rw','kode_pos','dusun','desa','kecamatan','kabupaten','provinsi','no_urut_kk','status_hubungan','nik','nama_lengkap','jenis_kelamin','tempat_lahir','tanggal_lahir','agama','pendidikan','jenis_pekerjaan','status_perkawinan','kewarganegaraan','no_paspor','no_kitas_kitap','nama_ayah','nama_ibu','golongan_darah']

def _pad3(value):
    if value is None:return None
    text=str(value).strip(); return text.zfill(3) if text.isdigit() and len(text)<=3 else text

def build_sid_workbook(scan_items):
    wb=Workbook(); ws=wb.active; ws.title='Data Kependudukan'; ws.append(COLUMNS); fill=PatternFill('solid',fgColor='D9EAF7')
    for cell in ws[1]: cell.font=Font(bold=True); cell.fill=fill; cell.alignment=Alignment(horizontal='center',vertical='center')
    row_count=0
    for item in scan_items:
        r=item.kk_record
        if not r: continue
        for m in r.members:
            ws.append([r.no_kk or '',r.nama_kepala_keluarga or '',r.alamat or '',_pad3(r.rt) or '',_pad3(r.rw) or '',r.kode_pos or '',r.dusun or '',r.desa or '',r.kecamatan or '',r.kabupaten or '',r.provinsi or '',m.no_urut_kk,m.status_hubungan or '',m.nik or '',m.nama_lengkap or '',m.jenis_kelamin or '',m.tempat_lahir or '',m.tanggal_lahir.strftime('%d-%m-%Y') if m.tanggal_lahir else '',m.agama or '',m.pendidikan or '',m.jenis_pekerjaan or '',m.status_perkawinan or '',m.kewarganegaraan or '',m.no_paspor or '',m.no_kitas_kitap or '',m.nama_ayah or '',m.nama_ibu or '',m.golongan_darah or '']); row_count+=1; er=ws.max_row
            for col in (1,4,5,14): ws.cell(er,col).number_format='@'
    ws.freeze_panes='A2'; ws.auto_filter.ref=f'A1:AB{ws.max_row}'
    out=BytesIO(); wb.save(out); return out.getvalue(),row_count
